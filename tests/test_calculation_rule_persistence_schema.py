from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook
import pytest
from sqlalchemy import delete, inspect, select
from sqlalchemy.exc import IntegrityError

from apps.api.app.database import Base
from apps.api.app.calculation_rules.comparison import CachedValueComparator
from apps.api.app.calculation_rules.compiler import FormulaCompiler
from apps.api.app.calculation_rules.evaluator import FormulaExecution, ScalarValue
from apps.api.app.calculation_rules.inventory import WorkbookFormulaInventory
from apps.api.app.calculation_rules.models import (
    CalculationRuleExtraction,
    ExecutableFormulaRule,
    FormulaCanonicalMapping,
    FormulaExecutionResultRecord,
    FormulaReferenceRecord,
    WorkbookFormulaCellRecord,
)
from apps.api.app.calculation_rules.repository import (
    CalculationRuleRepository,
    FormulaCanonicalMappingData,
    FormulaExecutionPersistenceData,
)
from apps.api.app.calculation_rules.types import (
    CalculationRuleExtractionConfiguration,
    FormulaIdFactory,
)
from apps.api.app.model_extraction_models import ModelVersion
from apps.api.app.model_extraction_repository import WorkbookVersionRepository
from apps.api.app.model_extraction_types import new_uuid
from apps.api.app.workbook_storage import DatabaseWorkbookStorage
from tests.model_extraction_test_support import (
    create_sqlite_session_factory,
    sqlite_file_url,
)


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    inputs["A1"] = 2
    inputs["A2"] = 3
    calc = workbook.create_sheet("Calc")
    calc["B1"] = "=Inputs!A1+1"
    calc["B2"] = '=COUNTIF(Inputs!A1:A2,">0")'
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.fixture
def persistence_context(tmp_path: Path):
    database_path = tmp_path / "rules.db"
    engine, session_factory = create_sqlite_session_factory(sqlite_file_url(database_path))
    Base.metadata.create_all(engine)
    session = session_factory()
    try:
        storage = DatabaseWorkbookStorage(session)
        workbook = WorkbookVersionRepository(session, storage).get_or_create(
            _workbook_bytes(), "rules.xlsx"
        )
        model = ModelVersion(
            id=new_uuid(),
            workbook_version_id=workbook.id,
            upload_filename="rules.xlsx",
            status="materialized",
            validation_status="validated",
            submitted=True,
            created_at=datetime.now(timezone.utc),
        )
        session.add(model)
        session.commit()
        yield engine, session_factory, session, workbook, model
    finally:
        session.close()
        engine.dispose()


def test_metadata_contains_exact_six_additive_rule_tables(persistence_context) -> None:
    engine, _factory, _session, _workbook, _model = persistence_context

    table_names = set(inspect(engine).get_table_names())

    assert {
        "calculation_rule_extractions",
        "workbook_formula_cells",
        "executable_formula_rules",
        "formula_references",
        "formula_canonical_mappings",
        "formula_execution_results",
    } <= table_names


def test_api_startup_registers_rule_metadata_without_importing_feature_first() -> None:
    script = (
        "from apps.api.app.main import Base; "
        "assert 'calculation_rule_extractions' in Base.metadata.tables"
    )

    completed = subprocess.run(
        [sys.executable, "-c", script],
        cwd=Path(__file__).parents[1],
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr


def test_schema_exposes_named_ownership_constraints_and_indexes(
    persistence_context,
) -> None:
    engine, _factory, _session, _workbook, _model = persistence_context
    inspector = inspect(engine)

    run_foreign_keys = {
        fk["name"]: fk for fk in inspector.get_foreign_keys("calculation_rule_extractions")
    }
    result_foreign_keys = {
        fk["name"]: fk for fk in inspector.get_foreign_keys("formula_execution_results")
    }
    run_unique_names = {
        constraint["name"]
        for constraint in inspector.get_unique_constraints("calculation_rule_extractions")
    }
    result_index_names = {
        index["name"] for index in inspector.get_indexes("formula_execution_results")
    }
    mapping_indexes = {
        index["name"]: index
        for index in inspector.get_indexes("formula_canonical_mappings")
    }

    assert run_foreign_keys["fk_calc_rule_extractions_workbook"]["options"]["ondelete"] == "RESTRICT"
    assert run_foreign_keys["fk_calc_rule_extractions_model"]["options"]["ondelete"] == "RESTRICT"
    assert result_foreign_keys["fk_formula_results_extraction"]["options"]["ondelete"] == "CASCADE"
    assert "uq_calc_rule_extractions_identity" in run_unique_names
    assert "ix_formula_results_extraction_status" in result_index_names
    assert mapping_indexes["uq_formula_canonical_mappings_output"]["unique"] == 1


def test_run_status_constraint_rejects_unknown_state(persistence_context) -> None:
    _engine, _factory, session, workbook, model = persistence_context
    configuration = CalculationRuleExtractionConfiguration()
    run_id = FormulaIdFactory.extraction_id(model.id, workbook.id, configuration)
    session.add(
        CalculationRuleExtraction(
            id=run_id,
            workbook_version_id=workbook.id,
            model_version_id=model.id,
            inventory_version=configuration.inventory_version,
            compiler_version=configuration.compiler_version,
            ir_version=configuration.ir_version,
            engine_version=configuration.engine_version,
            function_registry_version=configuration.function_registry_version,
            semantics_profile=configuration.semantics_profile,
            configuration_hash=configuration.configuration_hash,
            status="invented",
        )
    )

    with pytest.raises(IntegrityError):
        session.commit()
    session.rollback()


def test_repository_reuses_immutable_inventory_and_compilation_on_retry(
    persistence_context,
) -> None:
    _engine, _factory, session, workbook, model = persistence_context
    configuration = CalculationRuleExtractionConfiguration()
    catalog = WorkbookFormulaInventory(configuration).scan(
        _workbook_bytes(), workbook.id
    )
    compilations = tuple(
        FormulaCompiler(configuration).compile(formula, catalog)
        for formula in catalog.formulas
    )
    run_id = FormulaIdFactory.extraction_id(model.id, workbook.id, configuration)
    repository = CalculationRuleRepository(session)

    repository.start_run(run_id, model.id, workbook.id, configuration)
    repository.save_compilation(catalog, compilations, configuration)
    session.commit()
    repository.save_compilation(catalog, compilations, configuration)
    session.commit()

    assert len(session.scalars(select(WorkbookFormulaCellRecord)).all()) == 2
    assert len(session.scalars(select(ExecutableFormulaRule)).all()) == 2
    assert len(session.scalars(select(FormulaReferenceRecord)).all()) == 2
    run = session.get(CalculationRuleExtraction, run_id)
    assert run.status == "running"
    assert run.configuration_hash == configuration.configuration_hash


def test_repository_persists_and_reloads_maximum_valid_output_after_restart(
    persistence_context,
) -> None:
    _engine, session_factory, session, workbook, model = persistence_context
    configuration = CalculationRuleExtractionConfiguration()
    catalog = WorkbookFormulaInventory(configuration).scan(
        _workbook_bytes(), workbook.id
    )
    compilations = tuple(
        FormulaCompiler(configuration).compile(formula, catalog)
        for formula in catalog.formulas
    )
    run_id = FormulaIdFactory.extraction_id(model.id, workbook.id, configuration)
    repository = CalculationRuleRepository(session)
    repository.start_run(run_id, model.id, workbook.id, configuration)
    repository.save_compilation(catalog, compilations, configuration)
    supported = next(item for item in compilations if item.support_status == "supported")
    unsupported = next(item for item in compilations if item.support_status == "unsupported")
    supported_cell = next(cell for cell in catalog.formulas if cell.id == supported.formula_cell_id)
    unsupported_cell = next(cell for cell in catalog.formulas if cell.id == unsupported.formula_cell_id)
    comparator = CachedValueComparator(configuration)
    mappings = (
        FormulaCanonicalMappingData(
            id=FormulaIdFactory.mapping_id(run_id, supported_cell.id, None, "output"),
            calculation_rule_extraction_id=run_id,
            formula_cell_id=supported_cell.id,
            reference_id=None,
            mapping_role="output",
            mapping_status="unmapped",
            entity_kind=None,
            entity_id=None,
            financial_series_value_id=None,
            warnings=("canonical_mapping_missing",),
        ),
    )
    results = (
        FormulaExecutionPersistenceData(
            id=FormulaIdFactory.execution_result_id(run_id, supported_cell.id),
            calculation_rule_extraction_id=run_id,
            formula_cell_id=supported_cell.id,
            expression_id=supported.expression_id,
            execution=FormulaExecution("executed", ScalarValue.number(3), None, ()),
            comparison=comparator.compare(ScalarValue.number(3), None, "missing"),
        ),
        FormulaExecutionPersistenceData(
            id=FormulaIdFactory.execution_result_id(run_id, unsupported_cell.id),
            calculation_rule_extraction_id=run_id,
            formula_cell_id=unsupported_cell.id,
            expression_id=unsupported.expression_id,
            execution=FormulaExecution("not_executable", None, None, (), ("unsupported",)),
            comparison=comparator.compare(None, None, "missing"),
        ),
    )
    repository.replace_outputs(run_id, mappings, results, configuration)
    repository.complete_run(
        run_id,
        status="completed_with_warning",
        summary={"formula_cells_total": 2, "formula_cells_executed": 1},
        warnings=("unsupported_formula_cells",),
    )
    session.commit()
    session.close()

    restarted_session = session_factory()
    try:
        reloaded = CalculationRuleRepository(restarted_session).load_result(run_id)

        assert reloaded.calculation_rule_extraction_id == run_id
        assert reloaded.status == "completed_with_warning"
        assert reloaded.summary["formula_cells_total"] == 2
        assert set(reloaded.cells_by_address) == {"Calc!B1", "Calc!B2"}
        assert reloaded.cells_by_address["Calc!B1"].execution.value == ScalarValue.number(3)
        assert reloaded.cells_by_address["Calc!B2"].execution.status == "not_executable"
        assert reloaded.cells_by_address["Calc!B1"].mappings[0].mapping_status == "unmapped"
    finally:
        restarted_session.close()


def test_deleting_run_cascades_mappings_and_results_not_workbook_compilation(
    persistence_context,
) -> None:
    _engine, _factory, session, workbook, model = persistence_context
    configuration = CalculationRuleExtractionConfiguration()
    catalog = WorkbookFormulaInventory(configuration).scan(_workbook_bytes(), workbook.id)
    compilations = tuple(
        FormulaCompiler(configuration).compile(formula, catalog)
        for formula in catalog.formulas
    )
    run_id = FormulaIdFactory.extraction_id(model.id, workbook.id, configuration)
    repository = CalculationRuleRepository(session)
    repository.start_run(run_id, model.id, workbook.id, configuration)
    repository.save_compilation(catalog, compilations, configuration)
    supported = next(item for item in compilations if item.support_status == "supported")
    cell = next(item for item in catalog.formulas if item.id == supported.formula_cell_id)
    mapping = FormulaCanonicalMappingData(
        id=FormulaIdFactory.mapping_id(run_id, cell.id, None, "output"),
        calculation_rule_extraction_id=run_id,
        formula_cell_id=cell.id,
        reference_id=None,
        mapping_role="output",
        mapping_status="unmapped",
        entity_kind=None,
        entity_id=None,
        financial_series_value_id=None,
        warnings=(),
    )
    result = FormulaExecutionPersistenceData(
        id=FormulaIdFactory.execution_result_id(run_id, cell.id),
        calculation_rule_extraction_id=run_id,
        formula_cell_id=cell.id,
        expression_id=supported.expression_id,
        execution=FormulaExecution("executed", ScalarValue.number(3), None, ()),
        comparison=CachedValueComparator().compare(ScalarValue.number(3), None, "missing"),
    )
    repository.replace_outputs(run_id, (mapping,), (result,), configuration)
    session.commit()

    session.execute(delete(CalculationRuleExtraction).where(CalculationRuleExtraction.id == run_id))
    session.commit()

    assert session.scalar(select(FormulaCanonicalMapping.id)) is None
    assert session.scalar(select(FormulaExecutionResultRecord.id)) is None
    assert len(session.scalars(select(WorkbookFormulaCellRecord)).all()) == 2
    assert len(session.scalars(select(ExecutableFormulaRule)).all()) == 2


def test_alembic_upgrades_empty_sqlite_database_to_rule_extraction_head(
    tmp_path: Path,
) -> None:
    from alembic import command
    from alembic.config import Config

    database_path = tmp_path / "migration.db"
    config_path = Path(__file__).parents[1] / "apps" / "api" / "alembic.ini"
    config = Config(str(config_path))
    config.set_main_option("sqlalchemy.url", sqlite_file_url(database_path))

    command.upgrade(config, "head")

    engine, _factory = create_sqlite_session_factory(sqlite_file_url(database_path))
    try:
        assert {
            "calculation_rule_extractions",
            "workbook_formula_cells",
            "executable_formula_rules",
            "formula_references",
            "formula_canonical_mappings",
            "formula_execution_results",
        } <= set(inspect(engine).get_table_names())
    finally:
        engine.dispose()
